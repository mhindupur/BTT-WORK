import logging
import uuid
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

log = logging.getLogger(__name__)


def _norm(s: str | None) -> str:
    if s is None:
        return ""
    return str(s).strip().lower().replace(" ", "_")


def _cell(row: dict[str, Any], *names: str) -> Any:
    for n in names:
        k = _norm(n)
        for key, val in row.items():
            if _norm(key) == k:
                return val
    return None


def _dec(v: Any) -> Decimal | None:
    if v is None or v == "":
        return None
    try:
        return Decimal(str(v))
    except (InvalidOperation, ValueError):
        return None


def _int(v: Any) -> int | None:
    if v is None or v == "":
        return None
    try:
        return int(float(v))
    except (ValueError, TypeError):
        return None


def parse_payment_workbook(content: bytes) -> list[dict[str, Any]]:
    wb = load_workbook(filename=BytesIO(content), read_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if not header:
        return []
    keys = [str(h).strip() if h is not None else "" for h in header]
    out: list[dict[str, Any]] = []
    for r in rows_iter:
        if not r or all(c is None or str(c).strip() == "" for c in r):
            continue
        row = {keys[i]: r[i] for i in range(min(len(keys), len(r)))}
        out.append(
            {
                "vehicle_number": _cell(row, "vehicle_number", "vehicle no", "vehicle"),
                "driver_name": _cell(row, "driver_name", "driver name", "driver"),
                "driver_phone": _cell(row, "driver_phone", "mobile", "phone", "contact"),
                "trip_count": _int(_cell(row, "trip_count", "no_of_trips", "trips", "no of trips")),
                "fuel_advance": _dec(_cell(row, "fuel_advance", "fuel advance")),
                "emi": _dec(_cell(row, "emi", "load_emi", "load emi")),
                "other_advance": _dec(_cell(row, "other_advance", "other advance")),
                "net_payable": _dec(_cell(row, "net_payable", "net", "payable")),
                "raw": {k: v for k, v in row.items() if k},
            }
        )
    return out


def row_to_payment_line_dict(parsed: dict[str, Any]) -> dict[str, Any]:
    return {
        "vehicle_number": str(parsed["vehicle_number"]) if parsed.get("vehicle_number") else None,
        "driver_name": str(parsed["driver_name"]) if parsed.get("driver_name") else None,
        "driver_phone": str(parsed["driver_phone"]) if parsed.get("driver_phone") else None,
        "trip_count": parsed.get("trip_count"),
        "fuel_advance": parsed.get("fuel_advance"),
        "emi": parsed.get("emi"),
        "other_advance": parsed.get("other_advance"),
        "net_payable": parsed.get("net_payable"),
        "raw_row": parsed.get("raw"),
        "public_token": str(uuid.uuid4()),
    }
